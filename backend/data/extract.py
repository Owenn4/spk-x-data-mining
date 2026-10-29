"""
Ekstraksi data fundamental 4 variabel (Revenue, Net Income, Assets, Equity)
dari workbook laporan keuangan tahunan IDX (.xlsx) ke dataset tabular.
Menangani perbandingan Prior Year/Current Year serta penyesuaian satuan pelaporan.
"""

from __future__ import annotations

import argparse
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

# Suppress warning openpyxl terkait data validation extensions
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

PROJECT_ROOT = Path(__file__).resolve().parent
DATA_DIR = PROJECT_ROOT / "data"
RAW_DIR = DATA_DIR / "raw"
PROCESSED_DIR = DATA_DIR / "processed"
DEFAULT_OUTPUT_XLSX = PROCESSED_DIR / "financial_raw.xlsx"
DEFAULT_OUTPUT_CSV = PROCESSED_DIR / "financial_raw.csv"


def normalize(text: Any) -> str:

    """Normalisasi teks untuk pencocokan label tanpa terpengaruh kapitalisasi atau spasi ganda."""
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text).strip().lower())


def get_metadata(wb: openpyxl.Workbook, fallback_ticker: str) -> dict[str, str]:
    """Ekstraksi metadata emiten dari sheet [1000000] General information."""
    meta = {
        "ticker": fallback_ticker,
        "company": "",
        "rounding": "",
        "currency": "IDR",
        "unit": "IDR",
    }
    if "1000000" in wb.sheetnames:
        ws = wb["1000000"]
        for row in ws.iter_rows(values_only=True):
            if len(row) >= 2 and row[0]:
                lbl0 = normalize(str(row[0]))
                val = row[1]
                if lbl0 == "mata uang pelaporan":
                    val_norm = normalize(str(val))
                    if "usd" in val_norm or "dollar" in val_norm:
                        meta["currency"] = "USD"
                    elif "idr" in val_norm or "rupiah" in val_norm:
                        meta["currency"] = "IDR"
                elif "pembulatan yang digunakan" in lbl0 or lbl0 == "pembulatan":
                    meta["rounding"] = str(val).strip()

            if len(row) >= 3 and row[2]:
                lbl = normalize(str(row[2]))
                val = row[1]
                if lbl == "entity code" and val:
                    meta["ticker"] = str(val).strip().upper()
                elif lbl == "entity name" and val:
                    meta["company"] = str(val).strip()
                elif lbl == "description of presentation currency" and val:
                    val_norm = normalize(str(val))
                    if "usd" in val_norm or "dollar" in val_norm:
                        meta["currency"] = "USD"
                    elif "idr" in val_norm or "rupiah" in val_norm:
                        meta["currency"] = "IDR"
                elif lbl == "level of rounding used in financial statements" and val:
                    meta["rounding"] = str(val).strip()

    r_norm = normalize(meta["rounding"])
    curr = meta["currency"]
    if "billion" in r_norm or "miliar" in r_norm:
        meta["unit"] = f"{curr} billion"
    elif "million" in r_norm or "juta" in r_norm:
        meta["unit"] = f"{curr} million"
    elif "thousand" in r_norm or "ribu" in r_norm:
        meta["unit"] = f"{curr} thousand"
    else:
        meta["unit"] = curr

    return meta


def scale_val(val: float | None, unit: str, currency: str = "IDR") -> float | int | None:
    """
    Standardisasi angka agar seragam dengan satuan pelaporan emiten.
    Jika emiten bersatuan miliaran/jutaan/ribuan namun pada laporan tertentu IDX mencatat nilai penuh,
    maka dinormalisasi sesuai faktor skala.
    """
    if val is None:
        return None
    v = float(val)
    if "billion" in unit and abs(v) >= 1e11:
        v /= 1e9
    elif "million" in unit and abs(v) >= 1e11:
        v /= 1e6
    elif "thousand" in unit:
        # Jika pelaporan dalam ribuan tetapi tercatat satuan penuh
        # Untuk USD, angka penuh berada pada orde ~10^8 sampai 10^10
        if currency == "USD" and abs(v) >= 1e8:
            v /= 1e3
        elif currency != "USD" and abs(v) >= 1e11:
            v /= 1e3

    if v.is_integer():
        return int(v)
    return round(v, 2)


def get_exact_match_value(
    rows: list[tuple[Any, ...]],
    targets: list[str],
    col_idx: int,
) -> float | None:
    """Mencari nilai angka pada baris yang memiliki label persis (exact match)."""
    t_set = {normalize(t) for t in targets}
    for row in rows:
        labels = [normalize(c) for c in row if isinstance(c, str)]
        if any(l in t_set for l in labels):
            if col_idx < len(row):
                val = row[col_idx]
                if isinstance(val, (int, float)):
                    return float(val)
    return None


def get_contains_match_value(
    rows: list[tuple[Any, ...]],
    targets: list[str],
    col_idx: int,
    exclude: list[str] | None = None,
) -> float | None:
    """Mencari nilai angka dengan pola substring bersyarat (dengan opsi eksklusi kata kunci)."""
    t_set = [normalize(t) for t in targets]
    exc_set = [normalize(e) for e in (exclude or [])]
    for row in rows:
        labels = [normalize(c) for c in row if isinstance(c, str)]
        if any(any(e in l for e in exc_set) for l in labels):
            continue
        if any(any(t in l for t in t_set) for l in labels):
            if col_idx < len(row):
                val = row[col_idx]
                if isinstance(val, (int, float)):
                    return float(val)
    return None


# ---------------------------------------------------------------------------
# Single File Extraction
# ---------------------------------------------------------------------------


def extract_report_file(filepath: Path) -> tuple[dict[str, str], dict[str, dict[str, Any]]]:
    """
    Membaca satu file laporan tahunan IDX XLSX.
    Mengembalikan tuple (metadata, data) dengan key 'current' (col 1) dan 'prior' (col 2).
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    meta = get_metadata(wb, fallback_ticker=filepath.parent.name)

    # Identifikasi sheet Neraca / Balance Sheet (1210000 umum, 3210000 infrastruktur/jasa)
    bs_candidates = ["1210000", "3210000"]
    bs_rows = []
    for cand in bs_candidates:
        if cand in wb.sheetnames:
            rows = list(wb[cand].iter_rows(values_only=True))
            if any(isinstance(c, (int, float)) and c != 0 for r in rows for c in r):
                bs_rows = rows
                break
    if not bs_rows and bs_candidates[0] in wb.sheetnames:
        bs_rows = list(wb[bs_candidates[0]].iter_rows(values_only=True))

    # Identifikasi sheet Laba Rugi / Income Statement (1321000 by function, 1311000 by nature, 3312000 jasa)
    pl_candidates = ["1321000", "1311000", "3312000", "3311000"]
    pl_rows = []
    for cand in pl_candidates:
        if cand in wb.sheetnames:
            rows = list(wb[cand].iter_rows(values_only=True))
            if any(isinstance(c, (int, float)) and c != 0 for r in rows for c in r):
                pl_rows = rows
                break
    if not pl_rows and pl_candidates[0] in wb.sheetnames:
        pl_rows = list(wb[pl_candidates[0]].iter_rows(values_only=True))

    notes_rows = list(wb["1617000"].iter_rows(values_only=True)) if "1617000" in wb.sheetnames else []

    extracted: dict[str, dict[str, Any]] = {"current": {}, "prior": {}}

    for col_idx, key in [(1, "current"), (2, "prior")]:
        # Total aset
        assets = get_exact_match_value(bs_rows, ["jumlah aset", "total assets"], col_idx)

        # Total ekuitas
        equity = get_exact_match_value(bs_rows, ["jumlah ekuitas", "total equity"], col_idx)

        # Laba bersih bagian entitas induk
        ni = get_contains_match_value(
            pl_rows,
            ["entitas induk", "parent entity"],
            col_idx,
            exclude=["komprehensif", "comprehensive", "per saham", "earnings per share"],
        )
        if ni is None:
            ni = get_exact_match_value(pl_rows, ["jumlah laba (rugi)", "total profit (loss)"], col_idx)

        # Pendapatan usaha / penjualan bersih
        rev = get_exact_match_value(
            pl_rows,
            ["penjualan dan pendapatan usaha", "sales and revenue", "pendapatan", "revenue"],
            col_idx,
        )

        # Fallback sheet 1617000 jika revenue utama kosong (contoh ASII/UNTR 2023)
        if rev is None and notes_rows:
            ncol = 2 if key == "current" else 6
            rev = get_exact_match_value(notes_rows, ["tipe pendapatan", "type of revenue"], ncol)
            if rev is None:
                serv = get_exact_match_value(notes_rows, ["pendapatan dari jasa", "service revenue"], ncol)
                prod = get_exact_match_value(notes_rows, ["pendapatan dari produk", "product revenue"], ncol)
                if serv is not None or prod is not None:
                    rev = (serv or 0.0) + (prod or 0.0)

        curr_curr = meta.get("currency", "IDR")
        extracted[key] = {
            "revenue": scale_val(rev, meta["unit"], curr_curr),
            "net_income": scale_val(ni, meta["unit"], curr_curr),
            "total_assets": scale_val(assets, meta["unit"], curr_curr),
            "total_equity": scale_val(equity, meta["unit"], curr_curr),
        }

    wb.close()
    return meta, extracted


def extract_stock_series(
    ticker_dir: Path,
    target_data_years: list[int],
) -> list[dict[str, Any]]:
    """
    Mengekstrak rentang data_year (misal 2021..2025) untuk satu ticker.
    Metodologi:
    - 2021 diambil dari Prior Year file report 2022 (audited restated).
    - 2022..2025 diambil dari Current Year file masing-masing tahun.
    - Dilengkapi fallback cross-population antar-laporan jika terdapat baris kosong.
    """
    report_files = sorted(ticker_dir.glob("*Tahunan*.xlsx"))
    if not report_files:
        report_files = sorted(ticker_dir.glob("*.xlsx"))

    if not report_files:
        return []

    files_by_year: dict[int, Path] = {}
    for f in report_files:
        m = re.search(r"FinancialStatement-(\d{4})-Tahunan", f.name)
        if m:
            files_by_year[int(m.group(1))] = f

    if not files_by_year:
        return []

    # Parse file-file yang ada
    extracted_reports: dict[int, tuple[dict[str, str], dict[str, dict[str, Any]]]] = {}
    for ry, fpath in files_by_year.items():
        extracted_reports[ry] = extract_report_file(fpath)

    # Identifikasi informasi entitas
    ref_year = 2022 if 2022 in extracted_reports else sorted(extracted_reports.keys())[0]
    meta_ref = extracted_reports[ref_year][0]
    ticker = meta_ref["ticker"]
    company = meta_ref["company"]
    unit = meta_ref["unit"]

    series_rows: list[dict[str, Any]] = []

    for d_yr in target_data_years:
        if d_yr == 2021 and 2022 in extracted_reports:
            # 2021 diekstrak dari kolom Prior Year laporan 2022
            prior_data = extracted_reports[2022][1]["prior"]
            src_name = files_by_year[2022].name
            unit_yr = extracted_reports[2022][0]["unit"]
            series_rows.append(
                {
                    "ticker": ticker,
                    "company": company,
                    "data_year": 2021,
                    "report_year": 2022,
                    "revenue": prior_data["revenue"],
                    "net_income": prior_data["net_income"],
                    "total_assets": prior_data["total_assets"],
                    "total_equity": prior_data["total_equity"],
                    "unit": unit_yr,
                    "source_file": src_name,
                }
            )
        elif d_yr in extracted_reports:
            # Laporan tahunan d_yr tersedia
            cur_data = extracted_reports[d_yr][1]["current"]
            src_name = files_by_year[d_yr].name
            unit_yr = extracted_reports[d_yr][0]["unit"]

            # Fallback jika revenue kosong, cek kolom prior pada laporan tahun berikutnya
            rev = cur_data["revenue"]
            if rev is None and (d_yr + 1) in extracted_reports:
                next_prior = extracted_reports[d_yr + 1][1]["prior"]
                if next_prior.get("revenue") is not None:
                    rev = next_prior["revenue"]

            series_rows.append(
                {
                    "ticker": ticker,
                    "company": company,
                    "data_year": d_yr,
                    "report_year": d_yr,
                    "revenue": rev,
                    "net_income": cur_data["net_income"],
                    "total_assets": cur_data["total_assets"],
                    "total_equity": cur_data["total_equity"],
                    "unit": unit_yr,
                    "source_file": src_name,
                }
            )
        elif (d_yr + 1) in extracted_reports:
            # Jika laporan tahun d_yr tidak ada terpisah, ambil dari prior tahun berikutnya
            prior_data = extracted_reports[d_yr + 1][1]["prior"]
            src_name = files_by_year[d_yr + 1].name
            unit_yr = extracted_reports[d_yr + 1][0]["unit"]
            series_rows.append(
                {
                    "ticker": ticker,
                    "company": company,
                    "data_year": d_yr,
                    "report_year": d_yr + 1,
                    "revenue": prior_data["revenue"],
                    "net_income": prior_data["net_income"],
                    "total_assets": prior_data["total_assets"],
                    "total_equity": prior_data["total_equity"],
                    "unit": unit_yr,
                    "source_file": src_name,
                }
            )

    return series_rows


def run_extraction(
    raw_dir: Path,
    output_xlsx: Path,
    output_csv: Path | None = None,
    filter_stocks: list[str] | None = None,
    target_data_years: list[int] | None = None,
) -> pd.DataFrame:
    """Menjalankan ekstraksi menyeluruh ke seluruh folder di data/raw/."""
    if target_data_years is None:
        target_data_years = [2021, 2022, 2023, 2024, 2025]

    stock_dirs = sorted([d for d in raw_dir.iterdir() if d.is_dir()])
    if filter_stocks:
        f_upper = {s.upper().strip() for s in filter_stocks}
        stock_dirs = [d for d in stock_dirs if d.name.upper() in f_upper]

    print("=" * 80)
    print("EXTRACTION DATA LAPORAN KEUANGAN TAHUNAN IDX")
    print(f"Direktori sumber : {raw_dir}")
    print(f"Jumlah emiten    : {len(stock_dirs)} emiten ({', '.join(d.name for d in stock_dirs)})")
    print(f"Target data_year : {', '.join(str(y) for y in target_data_years)}")
    print("=" * 80)

    all_rows: list[dict[str, Any]] = []

    for sdir in stock_dirs:
        ticker = sdir.name
        print(f"Memproses emiten [{ticker}]...", end=" ", flush=True)
        try:
            rows = extract_stock_series(sdir, target_data_years)
            if rows:
                all_rows.extend(rows)
                print(f"OK ({len(rows)} periode diekstrak)")
            else:
                print("LEWAT (tidak ditemukan file laporan tahunan)")
        except Exception as e:
            print(f"GAGAL: {e}")

    columns = [
        "ticker",
        "company",
        "data_year",
        "report_year",
        "revenue",
        "net_income",
        "total_assets",
        "total_equity",
        "unit",
        "source_file",
    ]

    df = pd.DataFrame(all_rows, columns=columns)

    # Simpan output
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(output_xlsx, index=False, engine="openpyxl")
    print(f"\n[SUKSES] Dataset tersimpan ke Excel: {output_xlsx}")

    if output_csv:
        output_csv.parent.mkdir(parents=True, exist_ok=True)
        df.to_csv(output_csv, index=False, encoding="utf-8")
        print(f"[SUKSES] Dataset tersimpan ke CSV  : {output_csv}")

    # Tampilkan preview tabel hasil ekstraksi
    print("\n" + "=" * 80)
    print("PREVIEW DATASET KEUANGAN HASIL EKSTRAKSI")
    print("=" * 80)
    pd.set_option("display.max_columns", 15)
    pd.set_option("display.width", 1000)
    print(df.to_string(index=False))
    print("=" * 80 + "\n")

    return df


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Ekstraksi data laporan keuangan tahunan IDX ke format Excel & CSV.",
    )
    parser.add_argument(
        "--raw-dir",
        type=Path,
        default=RAW_DIR,
        help="Direktori tempat file laporan keuangan tahunan mentah disimpan (default: data/raw)",
    )
    parser.add_argument(
        "--output-xlsx",
        type=Path,
        default=DEFAULT_OUTPUT_XLSX,
        help="Path file output Excel (default: data/processed/financial_raw.xlsx)",
    )
    parser.add_argument(
        "--output-csv",
        type=Path,
        default=DEFAULT_OUTPUT_CSV,
        help="Path file output CSV (default: data/processed/financial_raw.csv)",
    )
    parser.add_argument(
        "--stocks",
        type=str,
        default=None,
        help="Daftar ticker yang akan diekstrak, dipisah koma (contoh: ASII,TLKM,UNTR). Default: semua folder.",
    )
    parser.add_argument(
        "--years",
        type=str,
        default="2021-2025",
        help="Rentang tahun data keuangan (contoh: 2021-2025). Default: 2021-2025.",
    )
    return parser.parse_args()


def parse_years_arg(years_str: str) -> list[int]:
    """Mengubah format string '2021-2025' atau '2021,2022' menjadi list integer tahun."""
    cleaned = years_str.strip()
    if "-" in cleaned:
        start_s, end_s = cleaned.split("-", 1)
        return list(range(int(start_s.strip()), int(end_s.strip()) + 1))
    return [int(y.strip()) for y in cleaned.split(",") if y.strip()]


def main() -> None:
    args = parse_args()
    filter_stocks = [s.strip().upper() for s in args.stocks.split(",")] if args.stocks else None
    target_years = parse_years_arg(args.years)

    run_extraction(
        raw_dir=args.raw_dir,
        output_xlsx=args.output_xlsx,
        output_csv=args.output_csv,
        filter_stocks=filter_stocks,
        target_data_years=target_years,
    )


if __name__ == "__main__":
    main()
